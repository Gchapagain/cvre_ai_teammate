"""Generate fictional sample files: a 12-unit industrial rent roll and a matching T-12.

Every name and number here is made up. Run:  python sample_data/generate_sample_data.py
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT_DIR = Path(__file__).resolve().parent
RENT_ROLL_PATH = OUT_DIR / "rent_roll.xlsx"
T12_PATH = OUT_DIR / "t12.xlsx"

PROPERTY_NAME = "Fallowmere Industrial Park (SAMPLE - fictional)"
AS_OF = date(2026, 9, 30)
OUTLIER_UNIT = "110"
OUTLIER_MULTIPLE = 2.3

# unit, tenant, SF, annual rent per SF (None = the outlier, worked out below), start, end
UNITS = [
    ("101", "Northfield Pallet Co.", 12000, 8.50, date(2022, 3, 1), date(2027, 2, 28)),
    ("102", "Brightline Logistics LLC", 18500, 8.25, date(2021, 7, 1), date(2028, 6, 30)),
    ("103", "Kestrel Machine Works", 9600, 9.10, date(2023, 1, 1), date(2027, 12, 31)),
    ("104", "Harbor & Pine Distribution", 22000, 7.95, date(2020, 10, 1), date(2030, 9, 30)),
    ("105", "Tallgrass Packaging Inc.", 15000, 8.75, date(2024, 5, 1), date(2029, 4, 30)),
    ("106", "Ironleaf Fabrication", 8400, 9.40, date(2023, 9, 1), date(2026, 12, 31)),
    ("107", "Meridian Cold Storage", 20000, 9.00, date(2022, 11, 1), date(2032, 10, 31)),
    ("108", "Quillon Auto Parts", 11200, 8.60, date(2025, 2, 1), date(2028, 1, 31)),
    ("109", "Sable Creek Brewing Supply", 7500, 9.25, date(2024, 8, 1), date(2027, 7, 31)),
    ("110", "Vantor Robotics Lab", 3200, None, date(2025, 6, 1), date(2030, 5, 31)),
    ("111", "Dunmore Freight Services", 16800, 8.10, date(2021, 4, 1), date(2027, 3, 31)),
    ("112", "Copperline Electric Supply", 10400, 8.90, date(2023, 6, 1), date(2028, 5, 31)),
]

RENT_ROLL_HEADERS = [
    "Unit", "Tenant", "SF", "Lease Start", "Lease End",
    "Monthly Base Rent", "Annual Base Rent", "Rent / SF (Annual)", "Lease Type",
]

T12_MONTHS = [date(2025 + (9 + i) // 12, (9 + i) % 12 + 1, 1) for i in range(12)]  # Oct 2025 - Sep 2026


def build_units():
    """Return unit rows with rents filled in, the outlier set to 2.3x the building average."""
    total_sf = sum(u[2] for u in UNITS)
    others_rent = sum(u[2] * u[3] for u in UNITS if u[3] is not None)
    outlier_sf = next(u[2] for u in UNITS if u[3] is None)
    # Building average = total annual rent / total SF, with the outlier included.
    outlier_rate = round(OUTLIER_MULTIPLE * others_rent / (total_sf - OUTLIER_MULTIPLE * outlier_sf), 2)

    rows = []
    for unit, tenant, sf, rate, start, end in UNITS:
        rate = outlier_rate if rate is None else rate
        monthly = round(sf * rate / 12, 2)
        annual = round(monthly * 12, 2)
        rows.append({
            "unit_number": unit,
            "tenant_name": tenant,
            "square_feet": sf,
            "lease_start": start,
            "lease_end": end,
            "monthly_rent": monthly,
            "annual_rent": annual,
            "rent_per_sf": round(annual / sf, 2),
            "lease_type": "NNN",
        })
    return rows


def write_rent_roll(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rent Roll"
    ws.append([PROPERTY_NAME])
    ws.append([f"Rent Roll as of {AS_OF.isoformat()}"])
    ws.append([])
    ws.append(RENT_ROLL_HEADERS)
    for r in rows:
        ws.append([
            r["unit_number"], r["tenant_name"], r["square_feet"], r["lease_start"], r["lease_end"],
            r["monthly_rent"], r["annual_rent"], r["rent_per_sf"], r["lease_type"],
        ])
    total_sf = sum(r["square_feet"] for r in rows)
    total_monthly = round(sum(r["monthly_rent"] for r in rows), 2)
    total_annual = round(sum(r["annual_rent"] for r in rows), 2)
    ws.append(["Total", None, total_sf, None, None, total_monthly, total_annual, round(total_annual / total_sf, 2), None])

    ws["A1"].font = Font(bold=True, size=14)
    for cell in ws[4] + ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        row[3].number_format = row[4].number_format = "yyyy-mm-dd"
        row[2].number_format = "#,##0"
        for i in (5, 6, 7):
            row[i].number_format = "#,##0.00"
    for col, width in zip("ABCDEFGHI", (8, 30, 10, 13, 13, 19, 18, 19, 11)):
        ws.column_dimensions[col].width = width
    wb.save(RENT_ROLL_PATH)


def write_t12(rows):
    base_rent = round(sum(r["monthly_rent"] for r in rows), 2)
    winter = {12, 1, 2}
    summer = {6, 7, 8}

    income = {
        "Base Rental Income": [base_rent] * 12,
        "CAM Reimbursements": [14250.00] * 12,
        "Other Income": [650.00 + 25 * (i % 3) for i in range(12)],
    }
    expenses = {
        "Property Taxes": [9800.00] * 12,
        "Insurance": [3150.00] * 12,
        "Repairs & Maintenance": [4200.00 + 350 * (i % 4) for i in range(12)],
        "Utilities": [2100.00 + (600 if m.month in winter | summer else 0) for m in T12_MONTHS],
        "Landscaping & Snow Removal": [1900.00 if m.month in winter else 950.00 for m in T12_MONTHS],
        "Management Fee": [round(0.03 * (base_rent + 14250.00), 2)] * 12,
    }

    def total(lines):
        return [round(sum(v[i] for v in lines.values()), 2) for i in range(12)]

    total_income, total_expenses = total(income), total(expenses)
    noi = [round(a - b, 2) for a, b in zip(total_income, total_expenses)]

    wb = Workbook()
    ws = wb.active
    ws.title = "T-12"
    ws.append([PROPERTY_NAME])
    ws.append([f"Trailing 12 Months Operating Statement, period ending {AS_OF.isoformat()}"])
    ws.append([])
    ws.append(["Line Item"] + [m.strftime("%b %Y") for m in T12_MONTHS] + ["Total"])

    def add(label, values, bold=False):
        ws.append([label] + values + [round(sum(values), 2)])
        if bold:
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)

    def section(title):
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    section("INCOME")
    for label, values in income.items():
        add(label, values)
    add("Total Income", total_income, bold=True)
    ws.append([])
    section("EXPENSES")
    for label, values in expenses.items():
        add(label, values)
    add("Total Expenses", total_expenses, bold=True)
    ws.append([])
    add("Net Operating Income", noi, bold=True)

    ws["A1"].font = Font(bold=True, size=14)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, min_col=2):
        for cell in row:
            cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 30
    for i in range(2, 15):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = 13
    wb.save(T12_PATH)


def main():
    rows = build_units()
    write_rent_roll(rows)
    write_t12(rows)

    total_sf = sum(r["square_feet"] for r in rows)
    average = sum(r["annual_rent"] for r in rows) / total_sf
    outlier = next(r for r in rows if r["unit_number"] == OUTLIER_UNIT)
    print(f"Wrote {RENT_ROLL_PATH.name} and {T12_PATH.name}")
    print(f"Building average rent/SF: {average:.2f}")
    print(f"Unit {OUTLIER_UNIT} rent/SF: {outlier['rent_per_sf']:.2f} ({outlier['rent_per_sf'] / average:.2f}x average)")


if __name__ == "__main__":
    main()
