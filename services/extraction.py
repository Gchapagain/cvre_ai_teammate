"""Extraction service: read a rent roll spreadsheet, have Gemini turn it into
strict JSON, validate that JSON, and save the rows to rent_roll_units.

Extraction only copies what the file says. Judging the numbers (for example the
rent/SF flag) is the validation service's job, so a row with a bad value (zero
SF, negative rent, text where a number belongs) is saved exactly as found.

The whole file is rejected, and nothing is saved, only when the file cannot be
read, a required column is missing, or Gemini's reply is not the JSON we asked for.
"""

import csv
import hashlib
import json
import os
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from google.genai import types
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, StrictFloat, StrictStr, ValidationError

import db
from services import gemini
from services.gemini import KEY_PLACEHOLDER  # noqa: F401  (used by the tests)


PROMPT = """You are extracting a commercial real estate rent roll into JSON.

Rules:
- Return one object per leased unit row in the table.
- Skip title rows, header rows, blank rows and total/subtotal rows.
- Copy every value exactly as it appears. Do not calculate, round, correct or guess.
- Dates must be YYYY-MM-DD. Numbers must be plain numbers with no currency symbols or commas.
- If a cell that should hold a number or a date holds something else, copy its text exactly.
- If a cell is empty, use null.
- unit_number is text, exactly as written.

Rent roll (one spreadsheet row per line, cells separated by " | "):

{sheet_text}
"""


class ExtractionError(Exception):
    """Raised when the rent roll cannot be extracted or fails validation."""


# A numeric cell as found: normally a number, but text is kept if that is what the file has.
NumberAsFound = StrictFloat | StrictStr | None

# Column headings a rent roll must have, as patterns matched against lower-case header cells.
REQUIRED_COLUMNS = {
    "unit": r"^(unit|suite)",
    "tenant": r"tenant",
    "square feet": r"^(sf|rsf|sq\.? ?ft|square f|area)",
    "lease start": r"start|commence",
    "lease end": r"end|expir",
    "monthly rent": r"month",
    "annual rent": r"^(annual|yearly)",
    "rent per SF": r"/ ?sf|per sf|psf",
}


class RentRollUnit(BaseModel):
    """One rent roll row as extracted. Every field must be present, but values are not judged here."""

    model_config = ConfigDict(extra="forbid")

    unit_number: str = Field(min_length=1)
    tenant_name: str | None
    square_feet: NumberAsFound
    lease_start: str | None
    lease_end: str | None
    monthly_rent: NumberAsFound
    annual_rent: NumberAsFound
    rent_per_sf: NumberAsFound
    lease_type: str | None


class RentRollExtraction(BaseModel):
    model_config = ConfigDict(extra="forbid")

    units: list[RentRollUnit] = Field(min_length=1)


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _read_csv_rows(path):
    try:
        with open(path, newline="", encoding="utf-8-sig") as handle:
            rows = [[cell.strip() for cell in row] for row in csv.reader(handle)]
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ExtractionError(f"{Path(path).name} could not be read as a CSV file. Nothing was saved.") from exc
    return [r for r in rows if any(r)]


def read_sheet_rows(xlsx_path):
    """Return the first sheet (or the CSV file) as a list of rows, each a list of cell strings."""
    if Path(xlsx_path).suffix.lower() == ".csv":
        return _read_csv_rows(xlsx_path)
    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as exc:
        raise ExtractionError(
            f"{Path(xlsx_path).name} could not be read as an Excel file. Nothing was saved."
        ) from exc
    try:
        ws = wb.worksheets[0]
        rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return [r for r in rows if any(r)]


def _missing_columns(row):
    cells = [cell.strip().lower() for cell in row]
    return [name for name, pattern in REQUIRED_COLUMNS.items()
            if not any(re.search(pattern, cell) for cell in cells)]


def check_required_columns(source_rows, filename):
    """Reject the file if its header row lacks a required column."""
    # The header row is the one that matches the most required columns.
    missing = min((_missing_columns(row) for row in source_rows), key=len)
    if missing:
        raise ExtractionError(
            f"{filename} is missing required column(s): {', '.join(missing)}. Nothing was saved."
        )


def validate_extraction(raw_json, source_rows):
    """Check the model's JSON has the right shape and only units from the file. Returns the units."""
    try:
        extraction = RentRollExtraction.model_validate_json(raw_json)
    except ValidationError as exc:
        raise ExtractionError(f"Gemini output did not match the required JSON schema:\n{exc}") from exc

    # A unit ID repeated in the file is kept for validation to flag. One that Gemini
    # invented, or repeated more often than the file does, is a bad reply.
    in_file = Counter(cell.strip() for row in source_rows for cell in row)
    returned = Counter(unit.unit_number for unit in extraction.units)
    for unit_number, count in returned.items():
        if in_file[unit_number] == 0:
            raise ExtractionError(f"Unit {unit_number} is not in the source file.")
        if count > in_file[unit_number]:
            raise ExtractionError(f"Unit {unit_number} was returned more often than it appears in the file.")
    return extraction.units


def _build_config():
    """Gemini settings for extraction. GEMINI_THINKING_LEVEL in .env is optional."""
    thinking_level = os.getenv("GEMINI_THINKING_LEVEL")
    return types.GenerateContentConfig(
        temperature=0,
        response_mime_type="application/json",
        response_json_schema=RentRollExtraction.model_json_schema(),
        thinking_config=types.ThinkingConfig(thinking_level=thinking_level) if thinking_level else None,
    )


def call_gemini(client, sheet_text, model=None):
    """Send the rent roll text to Gemini and return its raw JSON string."""
    try:
        client = client or gemini.get_client()
        return gemini.generate(client, PROMPT.format(sheet_text=sheet_text), _build_config(), model)
    except gemini.GeminiQuotaError as exc:
        raise ExtractionError(
            f"{exc} Nothing was saved. Until then, replay a recorded response or use a key with a higher quota."
        ) from exc
    except gemini.GeminiError as exc:
        raise ExtractionError(f"{exc} Nothing was saved.") from exc


def _prompt_fingerprint(sheet_text):
    """A short ID for the exact prompt, so a recording is only replayed for the prompt it answered."""
    return hashlib.sha256(PROMPT.format(sheet_text=sheet_text).encode()).hexdigest()


def save_recording(path, raw_json, sheet_text, model=None):
    """Save a real Gemini response to a file so it can be replayed without calling the API."""
    recording = {
        "recorded_at": datetime.now().isoformat(timespec="seconds"),
        "model": gemini.model_name(model),
        "prompt_sha256": _prompt_fingerprint(sheet_text),
        "response_text": raw_json,
    }
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(recording, indent=2))


def load_recording(path, sheet_text):
    """Return the recorded Gemini response, refusing one made for a different file or prompt."""
    path = Path(path)
    if not path.exists():
        raise ExtractionError(f"No recorded Gemini response at {path}. Record one with a live run first.")
    recording = json.loads(path.read_text())
    if recording.get("prompt_sha256") != _prompt_fingerprint(sheet_text):
        raise ExtractionError(
            f"The recording at {path} is out of date: the prompt or the source file has changed "
            "since it was made. Record a new one with a live run."
        )
    return recording["response_text"]


def save_units(conn, deal_id, document_id, units):
    """Replace this deal's rent roll rows with the newly extracted ones."""
    conn.execute("DELETE FROM rent_roll_units WHERE deal_id = ?", (deal_id,))
    conn.executemany(
        """INSERT INTO rent_roll_units
           (deal_id, document_id, unit_number, tenant_name, square_feet, lease_start,
            lease_end, monthly_rent, annual_rent, rent_per_sf, lease_type)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        [
            (deal_id, document_id, u.unit_number, u.tenant_name, u.square_feet,
             u.lease_start, u.lease_end, u.monthly_rent,
             u.annual_rent, u.rent_per_sf, u.lease_type)
            for u in units
        ],
    )


def _mark_document_read(conn, deal_id, path, source, document_id):
    """Record that the file was read: update its existing documents row, or add one."""
    if document_id is None:
        return conn.execute(
            """INSERT INTO documents (deal_id, filename, doc_type, file_path, source, status, used_for)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (deal_id, path.name, "Rent roll", str(path), source, "Read", "Rent roll"),
        ).lastrowid
    conn.execute("UPDATE documents SET doc_type = 'Rent roll', status = 'Read', used_for = 'Rent roll' WHERE id = ?",
                 (document_id,))
    return document_id


def extract_rent_roll(xlsx_path, deal_id, client=None, db_path=None, model=None, source="Upload",
                      record_to=None, replay_from=None, document_id=None):
    """Extract a rent roll file into rent_roll_units. Returns the number of units saved.

    replay_from: use the Gemini response saved in this file instead of calling the API.
    record_to: after a live call, save Gemini's response to this file.
    document_id: the file's existing documents row, if it was already saved as Received.
    """
    xlsx_path = Path(xlsx_path)
    source_rows = read_sheet_rows(xlsx_path)
    if not source_rows:
        raise ExtractionError(f"{xlsx_path.name} has no data.")
    check_required_columns(source_rows, xlsx_path.name)
    sheet_text = "\n".join(" | ".join(row) for row in source_rows)

    if replay_from:
        raw_json = load_recording(replay_from, sheet_text)
    else:
        raw_json = call_gemini(client, sheet_text, model)
        if record_to:
            save_recording(record_to, raw_json, sheet_text, model)
    units = validate_extraction(raw_json, source_rows)

    db.init_db(db_path)
    conn = db.get_connection(db_path)
    try:
        with conn:
            save_units(conn, deal_id, _mark_document_read(conn, deal_id, xlsx_path, source, document_id), units)
    finally:
        conn.close()
    return len(units)
