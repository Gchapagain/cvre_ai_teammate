"""Compare every extracted rent roll value to the source spreadsheet.

Run from the project root:  .venv/bin/python -m unittest discover -s tests -t . -v

TestLiveExtraction calls the real Gemini API and is skipped until a real key is
in .env. The other tests use a stand-in for Gemini so they run without a key.
"""

import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from dotenv import load_dotenv
from openpyxl import load_workbook

import db
from services import extraction

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RENT_ROLL = PROJECT_ROOT / "sample_data" / "rent_roll.xlsx"
FIELDS = [
    "unit_number", "tenant_name", "square_feet", "lease_start", "lease_end",
    "monthly_rent", "annual_rent", "rent_per_sf", "lease_type",
]
NUMERIC_FIELDS = {"square_feet", "monthly_rent", "annual_rent", "rent_per_sf"}

load_dotenv(PROJECT_ROOT / ".env")
HAS_REAL_KEY = os.getenv("GEMINI_API_KEY", "").strip() not in ("", extraction.KEY_PLACEHOLDER)


def read_source_units():
    """Read the rent roll straight from the Excel file, without any AI."""
    ws = load_workbook(RENT_ROLL, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_index = next(i for i, r in enumerate(rows) if r[0] == "Unit")
    units = []
    for row in rows[header_index + 1:]:
        if row[0] is None or row[0] == "Total":
            continue
        unit = dict(zip(FIELDS, row))
        unit["lease_start"] = unit["lease_start"].date().isoformat()
        unit["lease_end"] = unit["lease_end"].date().isoformat()
        units.append(unit)
    return units


def fake_client(units):
    """A stand-in Gemini client that replies with the given units as JSON."""
    reply = SimpleNamespace(text=json.dumps({"units": units}))
    return SimpleNamespace(models=SimpleNamespace(generate_content=lambda **kwargs: reply))


class ExtractionTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db_path = Path(self.tmp.name) / "test.db"
        db.init_db(self.db_path)
        conn = db.get_connection(self.db_path)
        with conn:
            self.deal_id = conn.execute(
                "INSERT INTO deals (name, property_type) VALUES ('Sample Deal', 'Industrial')"
            ).lastrowid
        conn.close()
        self.source = read_source_units()

    def saved_units(self):
        conn = db.get_connection(self.db_path)
        rows = conn.execute(
            f"SELECT {', '.join(FIELDS)} FROM rent_roll_units WHERE deal_id = ?", (self.deal_id,)
        ).fetchall()
        conn.close()
        return {row["unit_number"]: dict(row) for row in rows}

    def find_mismatches(self):
        """Return a list of every value in the database that differs from the source file."""
        saved = self.saved_units()
        mismatches = []
        if len(saved) != len(self.source):
            mismatches.append(f"unit count: source {len(self.source)}, extracted {len(saved)}")
        for expected in self.source:
            actual = saved.get(expected["unit_number"])
            if actual is None:
                mismatches.append(f"unit {expected['unit_number']}: missing")
                continue
            for field in FIELDS:
                if field in NUMERIC_FIELDS:
                    same = abs(actual[field] - expected[field]) < 0.005
                else:
                    same = actual[field] == expected[field]
                if not same:
                    mismatches.append(
                        f"unit {expected['unit_number']} {field}: source {expected[field]!r}, extracted {actual[field]!r}"
                    )
        return mismatches


class TestSampleData(ExtractionTestCase):
    def test_twelve_units_with_one_outlier_at_2_3x(self):
        self.assertEqual(len(self.source), 12)
        average = sum(u["annual_rent"] for u in self.source) / sum(u["square_feet"] for u in self.source)
        ratios = {u["unit_number"]: u["rent_per_sf"] / average for u in self.source}
        above_2x = [unit for unit, ratio in ratios.items() if ratio > 2]
        self.assertEqual(above_2x, ["110"])
        self.assertAlmostEqual(ratios["110"], 2.3, places=2)


class TestPipelineWithoutGemini(ExtractionTestCase):
    def test_every_value_saved_matches_source(self):
        count = extraction.extract_rent_roll(
            RENT_ROLL, self.deal_id, client=fake_client(self.source), db_path=self.db_path
        )
        self.assertEqual(count, 12)
        self.assertEqual(self.find_mismatches(), [])

    def test_comparison_catches_a_wrong_value(self):
        wrong = [dict(u) for u in self.source]
        wrong[4]["monthly_rent"] += 100
        extraction.extract_rent_roll(RENT_ROLL, self.deal_id, client=fake_client(wrong), db_path=self.db_path)
        mismatches = self.find_mismatches()
        self.assertEqual(len(mismatches), 1)
        self.assertIn("unit 105 monthly_rent", mismatches[0])

    def test_rerun_replaces_rows_instead_of_duplicating(self):
        for _ in range(2):
            extraction.extract_rent_roll(
                RENT_ROLL, self.deal_id, client=fake_client(self.source), db_path=self.db_path
            )
        self.assertEqual(len(self.saved_units()), 12)

    def test_bad_output_is_rejected_and_nothing_is_saved(self):
        missing_field = [{k: v for k, v in u.items() if k != "annual_rent"} for u in self.source]
        extra_field = [dict(u, notes="made up") for u in self.source]
        text_number = [dict(u, square_feet="twelve thousand") for u in self.source]
        invented_unit = self.source + [dict(self.source[0], unit_number="999")]
        duplicate_unit = self.source + [self.source[0]]
        cases = {
            "missing field": missing_field,
            "extra field": extra_field,
            "text where a number belongs": text_number,
            "unit not in the file": invented_unit,
            "duplicate unit": duplicate_unit,
            "no units": [],
        }
        for name, units in cases.items():
            with self.subTest(name):
                with self.assertRaises(extraction.ExtractionError):
                    extraction.extract_rent_roll(
                        RENT_ROLL, self.deal_id, client=fake_client(units), db_path=self.db_path
                    )
                self.assertEqual(self.saved_units(), {})

    def test_not_json_is_rejected(self):
        client = SimpleNamespace(
            models=SimpleNamespace(generate_content=lambda **kwargs: SimpleNamespace(text="Sure! Here it is"))
        )
        with self.assertRaises(extraction.ExtractionError):
            extraction.extract_rent_roll(RENT_ROLL, self.deal_id, client=client, db_path=self.db_path)


@unittest.skipUnless(HAS_REAL_KEY, "GEMINI_API_KEY in .env is still the placeholder")
class TestLiveExtraction(ExtractionTestCase):
    def test_gemini_extraction_matches_source_exactly(self):
        count = extraction.extract_rent_roll(RENT_ROLL, self.deal_id, db_path=self.db_path)
        self.assertEqual(count, 12)
        mismatches = self.find_mismatches()
        self.assertEqual(mismatches, [], "\n" + "\n".join(mismatches))
        print(f"\n  Gemini: {count} units x {len(FIELDS)} fields = {count * len(FIELDS)} values, all match the source")


if __name__ == "__main__":
    unittest.main()
