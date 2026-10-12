"""Tests for the download files: every exported value must equal what is saved for the deal."""

import io
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import db
import seed
from app import exports
from tests.test_extraction import read_source_units


class ExportTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.tmp.cleanup)
        cls.db_path = Path(cls.tmp.name) / "exports.db"
        deal_id = seed.seed(cls.db_path)["fallowmere"]
        conn = db.get_connection(cls.db_path)
        rows = lambda sql: [dict(row) for row in conn.execute(sql, (deal_id,))]
        cls.units = rows("SELECT * FROM rent_roll_units WHERE deal_id = ? ORDER BY id")
        cls.flags = rows("SELECT * FROM flags WHERE deal_id = ?")
        cls.t12 = rows("SELECT category, line_item, period, amount FROM t12_line_items WHERE deal_id = ? ORDER BY id")
        conn.close()


class TestRentRollExport(ExportTestCase):
    def setUp(self):
        self.table = exports.rent_roll_table(self.units, self.flags)

    def test_all_108_values_match_the_sample_file(self):
        source = read_source_units()
        checked = 0
        for expected, (_, row) in zip(source, self.table.iterrows(), strict=True):
            for field, column in exports.RENT_ROLL_COLUMNS.items():
                self.assertEqual(row[column], expected[field], f"unit {expected['unit_number']} {field}")
                checked += 1
        self.assertEqual(checked, 108)

    def test_the_flag_travels_with_the_export(self):
        flagged = self.table[self.table["Flag"] == "Flagged"]
        self.assertEqual(list(flagged["Unit"]), ["110"])
        self.assertIn("Unit 110 rent/SF is $20.33", flagged["Flag reason"].iloc[0])
        self.assertEqual(set(self.table[self.table["Unit"] != "110"]["Flag reason"]), {""})

    def test_csv_and_excel_hold_the_same_values_as_the_table(self):
        from_csv = pd.read_csv(io.BytesIO(exports.to_csv(self.table)), dtype={"Unit": str}, keep_default_na=False)
        from_excel = pd.read_excel(io.BytesIO(exports.to_excel(self.table, "Rent roll")), dtype={"Unit": str},
                                   keep_default_na=False)
        for copy in (from_csv, from_excel):
            self.assertEqual(list(copy.columns), list(self.table.columns))
            self.assertEqual(copy.values.tolist(), self.table.values.tolist())


class TestT12Export(ExportTestCase):
    def setUp(self):
        self.table = exports.t12_table(self.t12)

    def test_layout_is_one_row_per_line_item_in_statement_order(self):
        self.assertEqual(list(self.table.columns[:3]), ["Line item", "Category", "Oct 2025"])
        self.assertEqual(self.table.columns[-1], "Total")
        self.assertEqual(len(self.table.columns), 15)
        self.assertEqual(self.table["Line item"].iloc[0], "Base Rental Income")
        self.assertEqual(self.table["Line item"].iloc[-1], "Net Operating Income")
        self.assertEqual(len(self.table), 12)

    def test_all_156_values_match_the_sample_file(self):
        rows = list(load_workbook(seed.T12_FILE, data_only=True).worksheets[0].iter_rows(values_only=True))
        periods = rows[3][1:]
        file_values = {(row[0], period): amount for row in rows[4:] if row[0] and row[1] is not None
                       for period, amount in zip(periods, row[1:])}
        exported = {(row["Line item"], period): row[period]
                    for _, row in self.table.iterrows() for period in self.table.columns[2:]}
        self.assertEqual(len(exported), 156)
        self.assertEqual(exported, file_values)

    def test_excel_file_opens_with_the_named_sheet(self):
        workbook = load_workbook(io.BytesIO(exports.to_excel(self.table, "T-12")))
        self.assertEqual(workbook.sheetnames, ["T-12"])
        self.assertEqual(workbook["T-12"].max_row, 13)


class TestFileNames(unittest.TestCase):
    def test_slug(self):
        self.assertEqual(exports.slug("Fallowmere Industrial Park (SAMPLE - fictional)"), "fallowmere_industrial_park")
        self.assertEqual(exports.slug("A/B: Odd  Name!"), "a_b_odd_name")
        self.assertEqual(exports.slug("(only brackets)"), "deal")


if __name__ == "__main__":
    unittest.main()
