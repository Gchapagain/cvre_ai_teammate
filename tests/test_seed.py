"""Tests for seed.py: the database must hold exactly what the sample Excel files say. No Gemini calls."""

import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import db
import seed
from services import validation
from tests.test_extraction import FIELDS, read_source_units


class SeedTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.tmp.cleanup)
        cls.db_path = Path(cls.tmp.name) / "seeded.db"
        cls.deals = seed.seed(cls.db_path)

    def query(self, sql, params=()):
        conn = db.get_connection(self.db_path)
        rows = [dict(row) for row in conn.execute(sql, params)]
        conn.close()
        return rows


class TestSeededDeals(SeedTestCase):
    def test_three_deals_with_the_three_statuses(self):
        deals = self.query("SELECT name, status FROM deals ORDER BY id")
        self.assertEqual([d["status"] for d in deals], ["Needs Review", "Extracting", "Approved"])
        for deal in deals:
            self.assertIn("SAMPLE - fictional", deal["name"])

    def test_fallowmere_is_in_a_real_submarket_but_stays_marked_fictional(self):
        deal = self.query("SELECT name, submarket FROM deals WHERE id = ?", (self.deals["fallowmere"],))[0]
        self.assertEqual(deal["submarket"], "Alliance, Fort Worth, TX")
        self.assertIn("SAMPLE - fictional", deal["name"])

    def test_one_demo_user_per_role(self):
        users = self.query("SELECT email, full_name, role FROM users ORDER BY id")
        self.assertEqual([u["role"] for u in users], ["Agent", "Analyst", "Administrator"])
        for user in users:
            self.assertTrue(user["email"].endswith("@example.com"))
            # The header shows name and role side by side, so the name must not repeat the role.
            self.assertNotIn(user["role"], user["full_name"])

    def test_extracting_deal_has_a_document_and_no_numbers(self):
        deal_id = self.deals["quillhaven"]
        documents = self.query("SELECT status, used_for FROM documents WHERE deal_id = ?", (deal_id,))
        self.assertEqual(documents, [{"status": "Received", "used_for": None}])
        for table in ("rent_roll_units", "t12_line_items", "flags"):
            self.assertEqual(self.query(f"SELECT id FROM {table} WHERE deal_id = ?", (deal_id,)), [])

    def test_approved_deal_has_five_clean_units(self):
        deal_id = self.deals["brindlewick"]
        self.assertEqual(len(self.query("SELECT id FROM rent_roll_units WHERE deal_id = ?", (deal_id,))), 5)
        self.assertEqual(self.query("SELECT id FROM flags WHERE deal_id = ?", (deal_id,)), [])

    def test_documents_that_supplied_numbers_are_marked_read(self):
        statuses = self.query(
            "SELECT DISTINCT status FROM documents WHERE id IN "
            "(SELECT document_id FROM rent_roll_units UNION SELECT document_id FROM t12_line_items)"
        )
        self.assertEqual(statuses, [{"status": "Read"}])

    def test_every_number_belongs_to_a_document(self):
        for table in ("rent_roll_units", "t12_line_items"):
            self.assertEqual(self.query(f"SELECT id FROM {table} WHERE document_id IS NULL"), [])


class TestSeededRentRoll(SeedTestCase):
    def test_all_108_values_match_the_excel_file(self):
        source = read_source_units()
        saved = self.query(
            f"SELECT {', '.join(FIELDS)} FROM rent_roll_units WHERE deal_id = ? ORDER BY id",
            (self.deals["fallowmere"],),
        )
        checked = 0
        for expected, actual in zip(source, saved, strict=True):
            for field in FIELDS:
                self.assertEqual(actual[field], expected[field], f"unit {expected['unit_number']} {field}")
                checked += 1
        self.assertEqual(checked, 108)

    def test_validation_flags_exactly_unit_110(self):
        flags = self.query("SELECT deal_id, unit_number, rule FROM flags")
        self.assertEqual(flags, [{
            "deal_id": self.deals["fallowmere"], "unit_number": "110", "rule": validation.RULE_RENT_OUTLIER,
        }])


class TestSeededT12(SeedTestCase):
    def read_cells(self):
        """Every number in the T-12 file as {(line item, period): amount}, read without seed.py's code."""
        rows = list(load_workbook(seed.T12_FILE, data_only=True).worksheets[0].iter_rows(values_only=True))
        periods = rows[3][1:]
        return {
            (row[0], period): amount
            for row in rows[4:] if row[0] and row[1] is not None
            for period, amount in zip(periods, row[1:])
        }

    def test_all_156_values_match_the_excel_file(self):
        expected = self.read_cells()
        saved = self.query(
            "SELECT line_item, period, amount FROM t12_line_items WHERE deal_id = ?", (self.deals["fallowmere"],)
        )
        self.assertEqual(len(expected), 156)
        self.assertEqual({(r["line_item"], r["period"]): r["amount"] for r in saved}, expected)
        self.assertEqual(len(saved), 156)

    def test_lines_are_in_the_right_category(self):
        rows = self.query("SELECT DISTINCT line_item, category FROM t12_line_items")
        categories = {r["line_item"]: r["category"] for r in rows}
        self.assertEqual(categories["Base Rental Income"], "Income")
        self.assertEqual(categories["Property Taxes"], "Expense")
        self.assertEqual(categories["Management Fee"], "Expense")
        for line in ("Total Income", "Total Expenses", "Net Operating Income"):
            self.assertEqual(categories[line], "Total")
        self.assertEqual(len(categories), 12)

    def test_noi_is_income_minus_expenses(self):
        totals = {
            r["line_item"]: r["amount"]
            for r in self.query("SELECT line_item, amount FROM t12_line_items WHERE period = 'Total' AND category = 'Total'")
        }
        self.assertAlmostEqual(
            totals["Net Operating Income"], totals["Total Income"] - totals["Total Expenses"], places=2
        )


class TestResetFlag(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db_path = Path(self.tmp.name) / "local.db"

    def run_main(self, argv):
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            code = seed.main(argv, self.db_path)
        return code, output.getvalue()

    def test_without_reset_it_stops_and_deletes_nothing(self):
        self.db_path.write_bytes(b"pretend this is data worth keeping")
        code, output = self.run_main([])
        self.assertEqual(code, 1)
        self.assertIn("Nothing was changed", output)
        self.assertIn("--reset", output)
        self.assertEqual(self.db_path.read_bytes(), b"pretend this is data worth keeping")

    def test_without_reset_it_does_not_create_a_database(self):
        code, _ = self.run_main([])
        self.assertEqual(code, 1)
        self.assertFalse(self.db_path.exists())

    def test_with_reset_it_rebuilds_from_scratch(self):
        for _ in range(2):
            code, output = self.run_main(["--reset"])
            self.assertEqual(code, 0)
            self.assertIn("3 sample deals", output)
        conn = db.get_connection(self.db_path)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM deals").fetchone()[0], 3)
        conn.close()


if __name__ == "__main__":
    unittest.main()
