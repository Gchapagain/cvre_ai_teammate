"""Fill the local database with sample deals. No AI is used: the Excel files are read directly.

Run:  .venv/bin/python seed.py --reset

This deletes the local database and rebuilds it, so it only runs with --reset.
Every name and number is made up.
"""

import argparse
import json
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

import db
from services import validation

SAMPLE_DIR = Path(__file__).resolve().parent / "sample_data"
RENT_ROLL_FILE = SAMPLE_DIR / "rent_roll.xlsx"
T12_FILE = SAMPLE_DIR / "t12.xlsx"
DOCUMENT_SOURCE = "Sample data"
# A real industrial submarket, so market research has something to find. The property itself,
# its tenants and every number stay made up.
FALLOWMERE_MARKET = "Alliance, Fort Worth, TX"

UNIT_FIELDS = [
    "unit_number", "tenant_name", "square_feet", "lease_start", "lease_end",
    "monthly_rent", "annual_rent", "rent_per_sf", "lease_type",
]
T12_TOTAL_LINES = {"Total Income", "Total Expenses", "Net Operating Income"}
T12_SECTIONS = {"INCOME": "Income", "EXPENSES": "Expense"}

DEMO_USERS = [
    ("demo.broker@example.com", "Avery Sample", "Broker"),
]

# A second, clean rent roll for the Approved deal: unit, tenant, SF, annual rent per SF, start, end.
BRINDLEWICK_UNITS = [
    ("A", "Pemberly Sock Emporium", 1800, 26.00, "2023-04-01", "2028-03-31"),
    ("B", "Thistle & Fig Bakery", 1200, 28.50, "2024-02-01", "2029-01-31"),
    ("C", "Gantry Lane Books", 2400, 24.00, "2022-09-01", "2027-08-31"),
    ("D", "Odd Otter Toys", 1500, 27.00, "2025-01-01", "2030-12-31"),
    ("E", "Marrowby Tea Room", 3000, 25.50, "2021-06-01", "2031-05-31"),
]


def read_rent_roll(path):
    """Read unit rows from the sample rent roll, exactly as the file has them."""
    rows = list(load_workbook(path, data_only=True).worksheets[0].iter_rows(values_only=True))
    header = next(i for i, row in enumerate(rows) if row[0] == "Unit")
    units = []
    for row in rows[header + 1:]:
        if row[0] is None or row[0] == "Total":
            continue
        unit = dict(zip(UNIT_FIELDS, row))
        unit["lease_start"] = unit["lease_start"].date().isoformat()
        unit["lease_end"] = unit["lease_end"].date().isoformat()
        units.append(unit)
    return units


def read_t12(path):
    """Read the sample T-12 as (category, line item, period, amount) rows, one per cell."""
    rows = list(load_workbook(path, data_only=True).worksheets[0].iter_rows(values_only=True))
    header = next(i for i, row in enumerate(rows) if row[0] == "Line Item")
    periods = rows[header][1:]
    items, section = [], None
    for row in rows[header + 1:]:
        label = row[0]
        if label in T12_SECTIONS:
            section = T12_SECTIONS[label]
        elif label is not None:
            category = "Total" if label in T12_TOTAL_LINES else section
            items += [(category, label, period, amount) for period, amount in zip(periods, row[1:])]
    return items


def brindlewick_units():
    units = []
    for unit, tenant, sf, rate, start, end in BRINDLEWICK_UNITS:
        monthly = round(sf * rate / 12, 2)
        annual = round(monthly * 12, 2)
        units.append(dict(zip(UNIT_FIELDS, [unit, tenant, sf, start, end, monthly, annual, round(annual / sf, 2), "NNN"])))
    return units


def add_deal(conn, name, property_type, submarket, status):
    return conn.execute(
        "INSERT INTO deals (name, property_type, submarket, status) VALUES (?, ?, ?, ?)",
        (name, property_type, submarket, status),
    ).lastrowid


def add_document(conn, deal_id, filename, doc_type, status, used_for, file_path=None):
    return conn.execute(
        """INSERT INTO documents (deal_id, filename, doc_type, file_path, source, status, used_for)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        # Stored relative to the project folder, so the database works on any machine (and on Streamlit Cloud).
        (deal_id, filename, doc_type, file_path and str(Path(file_path).relative_to(db.PROJECT_ROOT)), DOCUMENT_SOURCE,
         status, used_for),
    ).lastrowid


def add_units(conn, deal_id, document_id, units):
    conn.executemany(
        f"INSERT INTO rent_roll_units (deal_id, document_id, {', '.join(UNIT_FIELDS)}) "
        f"VALUES (?, ?, {', '.join('?' for _ in UNIT_FIELDS)})",
        [(deal_id, document_id, *(unit[field] for field in UNIT_FIELDS)) for unit in units],
    )


def add_t12(conn, deal_id, document_id, items):
    conn.executemany(
        "INSERT INTO t12_line_items (deal_id, document_id, category, line_item, period, amount) VALUES (?, ?, ?, ?, ?, ?)",
        [(deal_id, document_id, *item) for item in items],
    )


SAMPLE_SOURCE = "Sample data"
SAMPLE_DRAFT_PREFIX = "SAMPLE draft"


def add_finished_steps(conn, deal_id):
    """The approved sample deal went through every step, so it has a draft and reviewed research.

    Both say plainly that they are made up for the demo. The draft has no figures (nothing to check),
    and the finding was dismissed, so it is never used in a narrative or sent to the AI.
    """
    conn.execute(
        "INSERT INTO drafts (deal_id, text, status) VALUES (?, ?, 'Approved')",
        (deal_id, f"{SAMPLE_DRAFT_PREFIX}, made up for the demo. This fictional retail row was reviewed and approved by the "
                  "demo broker. See the Review tab for the rent roll figures as extracted."))
    sources = json.dumps([{"name": "Sample data (fictional, no real source)", "url": None, "tier": None,
                           "date": None, "figure": None, "metric": None}])
    conn.execute(
        "INSERT INTO research_findings (deal_id, topic, summary, sources, status, connection_id) "
        "VALUES (?, 'Vacancy', ?, ?, 'dismissed', ?)",
        (deal_id, "SAMPLE finding, made up for the demo: a reviewed research card for a fictional submarket. "
                  "It was dismissed, so it is not used anywhere.", sources, SAMPLE_SOURCE))


def _anchor(days, clock):
    """A sample time: days after SAMPLE_ANCHOR_DAY at a clock time. Home shows the anchor day as today."""
    day = datetime.fromisoformat(db.SAMPLE_ANCHOR_DAY) + timedelta(days=days)
    return f"{day:%Y-%m-%d} {clock}:00"


def add_sample_inbox(conn, deals):
    """Made-up meetings and emails for the Home feed, all with source Sample (shown with a Sample chip).

    The people are made up and the deals are the fictional sample deals. No real inbox or calendar.
    """
    meetings = [
        (deals["fallowmere"], "Client call", "Dana Reyes", _anchor(0, "15:00"), _anchor(0, "15:30"), "Phone",
         "Walk through the flagged unit before pricing."),
        (None, "Coffee with a prospective client", "Jordan Okafor", _anchor(0, "08:30"), _anchor(0, "09:00"),
         "Lobby cafe", None),
        (deals["quillhaven"], "Site walk", "Marcus Bell", _anchor(1, "10:00"), _anchor(1, "11:00"), "On site", None),
        (deals["fallowmere"], "Pricing review", "Dana Reyes", _anchor(3, "14:00"), _anchor(3, "15:00"),
         "https://meet.example.com/sample-pricing-review", None),
        (deals["brindlewick"], "Owner check-in", "Ellis Park", _anchor(-1, "16:00"), _anchor(-1, "16:30"), "Phone", None),
    ]
    conn.executemany(
        """INSERT INTO meetings (deal_id, title, client_name, starts_at, ends_at, location, notes, source)
           VALUES (?, ?, ?, ?, ?, ?, ?, 'Sample')""", meetings)
    emails = [
        (deals["fallowmere"], "Dana Reyes", "Updated rent roll for Fallowmere",
         "Hi Avery, attached is the updated rent roll. Unit 110 is the one we discussed.", _anchor(0, "08:15"),
         ["fallowmere_rent_roll_v2.xlsx"]),
        (None, "Priya Nandakumar", "Flex space inquiry",
         "We are looking for about 20,000 SF of flex space. Requirements attached.", _anchor(0, "07:05"),
         ["requirements_sheet.pdf"]),
        (deals["quillhaven"], "Marcus Bell", "Quillhaven T-12 and lease abstract",
         "Here are the T-12 and the lease abstract you asked for.", _anchor(-1, "16:40"),
         ["quillhaven_t12.xlsx", "quillhaven_lease_abstract.pdf"]),
    ]
    conn.executemany(
        """INSERT INTO emails (deal_id, sender, subject, preview, received_at, attachment_names, source)
           VALUES (?, ?, ?, ?, ?, ?, 'Sample')""",
        [(deal, sender, subject, preview, received, json.dumps(files))
         for deal, sender, subject, preview, received, files in emails])


def seed(db_path=None):
    """Delete the database at db_path and rebuild it from sample data. Returns the deal IDs."""
    path = Path(db_path or db.DB_PATH)
    path.unlink(missing_ok=True)
    db.init_db(path)
    conn = db.get_connection(path)
    try:
        with conn:
            conn.executemany("INSERT INTO users (email, full_name, role) VALUES (?, ?, ?)", DEMO_USERS)

            fallowmere = add_deal(conn, "Fallowmere Industrial Park (SAMPLE - fictional)",
                                  "Industrial", FALLOWMERE_MARKET, "Needs Review")
            doc = add_document(conn, fallowmere, RENT_ROLL_FILE.name, "Rent roll", "Read", "Rent roll", RENT_ROLL_FILE)
            add_units(conn, fallowmere, doc, read_rent_roll(RENT_ROLL_FILE))
            doc = add_document(conn, fallowmere, T12_FILE.name, "T-12", "Read", "T-12", T12_FILE)
            add_t12(conn, fallowmere, doc, read_t12(T12_FILE))

            quillhaven = add_deal(conn, "Quillhaven Flex Center (SAMPLE - fictional)",
                                  "Flex", "Eastwick Bend (fictional)", "Extracting")
            add_document(conn, quillhaven, "quillhaven_rent_roll.xlsx", "Rent roll",
                         "Received", None)

            brindlewick = add_deal(conn, "Brindlewick Retail Row (SAMPLE - fictional)",
                                   "Retail", "Old Harrow Lane (fictional)", "Approved")
            doc = add_document(conn, brindlewick, "brindlewick_rent_roll.xlsx", "Rent roll", "Read", "Rent roll")
            add_units(conn, brindlewick, doc, brindlewick_units())
            # The sample Approved deal records who approved it, like any approval would.
            conn.execute("UPDATE deals SET approved_by = ?, approved_at = CURRENT_TIMESTAMP WHERE id = ?",
                         (DEMO_USERS[0][1], brindlewick))
            add_finished_steps(conn, brindlewick)
    finally:
        conn.close()

    deal_ids = {"fallowmere": fallowmere, "quillhaven": quillhaven, "brindlewick": brindlewick}
    conn = db.get_connection(path)
    try:
        with conn:
            add_sample_inbox(conn, deal_ids)
    finally:
        conn.close()
    for deal_id in deal_ids.values():
        validation.validate_deal(deal_id, path)
    return deal_ids


def main(argv=None, db_path=None):
    parser = argparse.ArgumentParser(description="Rebuild the local database from sample data.")
    parser.add_argument("--reset", action="store_true", help="delete the local database and rebuild it")
    args = parser.parse_args(argv)
    path = Path(db_path or db.DB_PATH)
    if not args.reset:
        print(
            "Nothing was changed.\n"
            f"seed.py deletes the local database ({path.name}) and rebuilds it from sample data.\n"
            "To go ahead, run:  .venv/bin/python seed.py --reset"
        )
        return 1
    deal_ids = seed(path)
    print(f"Rebuilt {path.name} at {datetime.now():%Y-%m-%d %H:%M} with {len(deal_ids)} sample deals.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
